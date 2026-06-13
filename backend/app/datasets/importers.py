from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import StringIO
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.db.models.document import SourceFile
from app.db.repositories.dataset import OpenDatasetRepository
from app.db.repositories.source_file import SourceFileRepository
from app.files.storage import LocalFileStorage


DATASET_COLUMN_MAP = {
    "dataset_name": "name",
    "description": "description",
    "category": "category",
    "department": "department",
    "owner": "owner",
    "source_platform": "source_platform",
    "source_url": "source_url",
    "open_type": "open_type",
    "update_frequency": "update_frequency",
    "update_date": "update_date",
    "time_range": "time_range",
    "region_level": "region_level",
    "sensitivity_level": "sensitivity_level",
    "download_count": "download_count",
    "api_count": "api_count",
}
FIELD_COLUMN_MAP = {
    "field_name",
    "field_type",
    "description",
    "example_value",
    "is_dimension",
    "is_measure",
}
SUPPORTED_TABULAR_TYPES = {"csv", "xlsx"}


class DatasetImportError(ValueError):
    pass


@dataclass(frozen=True)
class DatasetImportResult:
    catalog_source_file_id: UUID
    fields_source_file_id: UUID | None
    catalog_rows: int
    field_rows: int
    datasets_created: int
    datasets_updated: int
    fields_created: int
    fields_updated: int
    field_rows_skipped: int


class OpenDatasetImporter:
    def __init__(self, db: Session, storage: LocalFileStorage) -> None:
        self.db = db
        self.storage = storage
        self.source_files = SourceFileRepository(db)
        self.datasets = OpenDatasetRepository(db)

    def import_from_source_files(
        self,
        *,
        catalog_source_file_id: UUID,
        fields_source_file_id: UUID | None = None,
    ) -> DatasetImportResult:
        catalog_file = self._get_tabular_source_file(catalog_source_file_id, label="catalog")
        fields_file = (
            self._get_tabular_source_file(fields_source_file_id, label="fields")
            if fields_source_file_id is not None
            else None
        )

        try:
            catalog_rows = read_tabular_rows(
                self.storage.resolve(catalog_file.storage_uri),
                file_type=catalog_file.file_type,
            )
            field_rows = (
                read_tabular_rows(
                    self.storage.resolve(fields_file.storage_uri),
                    file_type=fields_file.file_type,
                )
                if fields_file is not None
                else []
            )
            result = self._import_rows(
                catalog_file=catalog_file,
                fields_file=fields_file,
                catalog_rows=catalog_rows,
                field_rows=field_rows,
            )
        except Exception as exc:
            catalog_file.parse_status = "failed"
            catalog_file.error_message = str(exc)
            if fields_file is not None:
                fields_file.parse_status = "failed"
                fields_file.error_message = str(exc)
            self.db.flush()
            raise

        catalog_file.parse_status = "parsed"
        catalog_file.error_message = None
        if fields_file is not None:
            fields_file.parse_status = "parsed"
            fields_file.error_message = None
        self.db.flush()
        return result

    def _get_tabular_source_file(self, source_file_id: UUID, *, label: str) -> SourceFile:
        source_file = self.source_files.get_by_id(source_file_id)
        if source_file is None:
            raise DatasetImportError(f"{label} source file not found")
        if source_file.file_type not in SUPPORTED_TABULAR_TYPES:
            raise DatasetImportError(
                f"{label} source file must be csv or xlsx, got {source_file.file_type}"
            )
        return source_file

    def _import_rows(
        self,
        *,
        catalog_file: SourceFile,
        fields_file: SourceFile | None,
        catalog_rows: list[dict[str, Any]],
        field_rows: list[dict[str, Any]],
    ) -> DatasetImportResult:
        if not catalog_rows:
            raise DatasetImportError("Dataset catalog is empty")

        now = datetime.now(UTC)
        dataset_by_name: dict[str, UUID] = {}
        datasets_created = 0
        datasets_updated = 0
        for row in catalog_rows:
            values = dataset_values_from_row(row, source_file=catalog_file, synced_at=now)
            dataset, created = self.datasets.upsert_dataset(values)
            dataset_by_name[dataset.name] = dataset.id
            if created:
                datasets_created += 1
            else:
                datasets_updated += 1

        fields_created = 0
        fields_updated = 0
        field_rows_skipped = 0
        for row in field_rows:
            dataset_name = clean_text(row.get("dataset_name"))
            field_name = clean_text(row.get("field_name"))
            if not dataset_name or not field_name:
                field_rows_skipped += 1
                continue
            dataset_id = dataset_by_name.get(dataset_name)
            if dataset_id is None:
                field_rows_skipped += 1
                continue
            _, created = self.datasets.upsert_field(
                field_values_from_row(row, dataset_id=dataset_id, source_file=fields_file)
            )
            if created:
                fields_created += 1
            else:
                fields_updated += 1

        return DatasetImportResult(
            catalog_source_file_id=catalog_file.id,
            fields_source_file_id=fields_file.id if fields_file is not None else None,
            catalog_rows=len(catalog_rows),
            field_rows=len(field_rows),
            datasets_created=datasets_created,
            datasets_updated=datasets_updated,
            fields_created=fields_created,
            fields_updated=fields_updated,
            field_rows_skipped=field_rows_skipped,
        )


def dataset_values_from_row(
    row: dict[str, Any],
    *,
    source_file: SourceFile,
    synced_at: datetime,
) -> dict[str, Any]:
    name = clean_text(row.get("dataset_name"))
    if not name:
        raise DatasetImportError("dataset_name is required")

    metadata = metadata_from_row(row, known_columns=set(DATASET_COLUMN_MAP))
    metadata["import_source_file_id"] = str(source_file.id)
    values = {
        "name": name,
        "description": clean_text(row.get("description")),
        "category": normalize_label(row.get("category")),
        "department": normalize_label(row.get("department")),
        "owner": clean_text(row.get("owner")),
        "source_platform": clean_text(row.get("source_platform")),
        "source_url": clean_text(row.get("source_url")),
        "open_type": clean_text(row.get("open_type")),
        "update_frequency": clean_text(row.get("update_frequency")),
        "update_date": parse_date(row.get("update_date")),
        "time_range": clean_text(row.get("time_range")),
        "region_level": clean_text(row.get("region_level")),
        "sensitivity_level": clean_text(row.get("sensitivity_level")),
        "download_count": parse_int(row.get("download_count")),
        "api_count": parse_int(row.get("api_count")),
        "status": "active",
        "last_synced_at": synced_at,
        "metadata_json": metadata,
    }
    return values


def field_values_from_row(
    row: dict[str, Any],
    *,
    dataset_id: UUID,
    source_file: SourceFile | None,
) -> dict[str, Any]:
    _ = source_file
    return {
        "dataset_id": dataset_id,
        "field_name": clean_text(row.get("field_name")),
        "field_type": clean_text(row.get("field_type")),
        "description": clean_text(row.get("description")),
        "example_value": clean_text(row.get("example_value")),
        "is_dimension": parse_bool(row.get("is_dimension")),
        "is_measure": parse_bool(row.get("is_measure")),
    }


def read_tabular_rows(path: Path, *, file_type: str) -> list[dict[str, Any]]:
    if file_type == "csv":
        return read_csv_rows(path)
    if file_type == "xlsx":
        return read_xlsx_rows(path)
    raise DatasetImportError(f"Unsupported tabular file type: {file_type}")


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    raw = path.read_bytes()
    text = decode_text(raw)
    reader = csv.DictReader(StringIO(text))
    return [normalize_row(row) for row in reader]


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = [clean_header(value) for value in next(rows)]
    except StopIteration:
        return []
    result: list[dict[str, Any]] = []
    for values in rows:
        row = {header: normalize_cell(value) for header, value in zip(headers, values, strict=False)}
        if any(value is not None for value in row.values()):
            result.append(row)
    return result


def normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    return {clean_header(key): normalize_cell(value) for key, value in row.items() if key is not None}


def metadata_from_row(row: dict[str, Any], *, known_columns: set[str]) -> dict[str, Any]:
    return {
        key: value
        for key, value in row.items()
        if key not in known_columns and value not in (None, "")
    }


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def clean_header(value: Any) -> str:
    return str(value or "").strip()


def normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def clean_text(value: Any) -> str | None:
    value = normalize_cell(value)
    if value is None:
        return None
    return str(value).strip() or None


def normalize_label(value: Any) -> str | None:
    text = clean_text(value)
    if text is None:
        return None
    return " ".join(text.split())


def parse_date(value: Any) -> date | None:
    value = normalize_cell(value)
    if value is None:
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    text = str(value).strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError as exc:
        raise DatasetImportError(f"Invalid date value: {value}") from exc


def parse_int(value: Any) -> int:
    value = normalize_cell(value)
    if value in (None, ""):
        return 0
    try:
        return int(float(str(value).replace(",", "")))
    except ValueError as exc:
        raise DatasetImportError(f"Invalid integer value: {value}") from exc


def parse_bool(value: Any) -> bool:
    value = normalize_cell(value)
    if value in (None, ""):
        return False
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"true", "1", "yes", "y", "是"}:
        return True
    if normalized in {"false", "0", "no", "n", "否"}:
        return False
    raise DatasetImportError(f"Invalid boolean value: {value}")
